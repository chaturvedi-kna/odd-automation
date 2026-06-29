import logging
from pathlib import Path
from sqlalchemy import select, and_
from sqlalchemy.orm import Session

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.models.entry_instance_status import EntryInstanceStatus
from app.models.entry_instance_detail import EntryInstanceDetail
from app.models.entry_instance_relationship import EntryInstanceRelationship  # Added to resolve Issue 1
from app.models.prr_dump_row import PrrDumpRow
from app.models.rbar_dump_row import RbarDumpRow
from app.models.dump_snapshot import DumpSnapshot
from app.models.audit_log import AuditLog
from app.models.enums import DecisionType, ImplStatus

logger = logging.getLogger(__name__)

# Master Formatting Typography
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")

MASTER_DUMP_FONT = Font(bold=False, color="000000")
MASTER_ADD_FONT = Font(bold=True, color="9C0006")
MASTER_ADD_FILL = PatternFill("solid", fgColor="FFC7CE")
MASTER_DEL_FONT = Font(bold=True, color="9C6500")
MASTER_DEL_FILL = PatternFill("solid", fgColor="FFEB9C")

ADD_DECISIONS = {DecisionType.ADD.value, DecisionType.DEPENDENCY_ADD.value}
DEL_DECISIONS = {
    DecisionType.DELETE.value, 
    DecisionType.DEPENDENCY_DELETE.value,
    DecisionType.SUPERSEDE.value,
    DecisionType.SUPERSEDE_PENDING.value
}

PRR_DUMP_HEADERS = [
    "#Application Name", "Screen Name", "name", "priority", "param_1", "condOperator_1", "value_1",
    "param_2", "condOperator_2", "value_2", "param_3", "condOperator_3", "value_3", "param_4",
    "condOperator_4", "value_4", "param_5", "condOperator_5", "value_5", "param_6", "condOperator_6",
    "value_6", "action", "routeListName", "diamAnsCode", "errorMessage", "msgPriority", "msgCpyCfgSet",
    "vendorId", "targetPrtName", "peerRouteTable"
]

RBAR_DUMP_HEADERS = [
    "#Application Name", "Screen Name", "tableName", "startAddr", "endAddr",
    "destination", "pfxLength", "oldTableName", "oldStartAddr", "oldPfxLength"
]


def _ist_format(dt) -> str:
    if not dt:
        return ""
    from zoneinfo import ZoneInfo
    return dt.astimezone(ZoneInfo("Asia/Kolkata")).strftime("%Y-%m-%d %H:%M:%S")


def _generate_unique_sheet(wb, target_name: str):
    """
    Issue 9 Fix: Removed arbitrary string-stripping methods. 
    Strictly truncates up to the 31-character limit and handles collisions with clear suffixes.
    """
    truncated_base = target_name[:28]
    candidate = target_name[:31]
    counter = 1
    
    while candidate in wb.sheetnames:
        suffix = f"_{counter}"
        available_len = 31 - len(suffix)
        candidate = f"{truncated_base[:available_len]}{suffix}"
        counter += 1
        
    return wb.create_sheet(candidate)


def _style_header(ws, headers: list[str]):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def _auto_width(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)


# Issue 7 Fix: Cleaned unused 's' parameter definitions from data extractors
def _unpack_prr_vals(d: EntryInstanceDetail) -> list:
    p = d.raw_payload or {}
    return [
        p.get("Application Name", "Diameter"), p.get("Screen Name", "PeerRouteRule"),
        d.final_prt_rule or p.get("name", ""), p.get("priority", ""),
        p.get("param_1", ""), p.get("condOperator_1", ""), p.get("value_1", ""),
        p.get("param_2", ""), p.get("condOperator_2", ""), p.get("value_2", ""),
        p.get("param_3", ""), p.get("condOperator_3", ""), p.get("value_3", ""),
        p.get("param_4", ""), p.get("condOperator_4", ""), p.get("value_4", ""),
        p.get("param_5", ""), p.get("condOperator_5", ""), p.get("value_5", ""),
        p.get("param_6", ""), p.get("condOperator_6", ""), p.get("value_6", ""),
        p.get("action", "RouteToPrt"), p.get("routeListName", ""), p.get("diamAnsCode", ""),
        p.get("errorMessage", ""), p.get("msgPriority", ""), p.get("msgCpyCfgSet", ""),
        p.get("vendorId", ""), p.get("targetPrtName", ""), p.get("peerRouteTable", "")
    ]


# Issue 7 & 8 Fix: Cleaned parameters and removed the hardcoded "IMSI" fallback string
def _unpack_rbar_vals(d: EntryInstanceDetail) -> list:
    p = d.raw_payload or {}
    return [
        p.get("Application Name", "Rbar"), p.get("Screen Name", "AddressRange"),
        p.get("tableName") or None,  # Falls back cleanly to database null state instead of hardcoding
        int(d.start_addr) if d.start_addr is not None else None,
        int(d.end_addr) if d.end_addr is not None else None,
        d.destination or "", p.get("pfxLength", ""), p.get("oldTableName", ""),
        int(p["oldStartAddr"]) if p.get("oldStartAddr") not in ("", None) else None,
        p.get("oldPfxLength", "")
    ]


def delta_excel(db: Session, request_id: str, out_path: str) -> str:
    """
    Generate a Delta Excel file for the given request.
    """
    rows = db.execute(
        select(EntryInstanceStatus, EntryInstanceDetail)
        .join(EntryInstanceDetail, EntryInstanceDetail.instance_status_id == EntryInstanceStatus.id)
        .where(EntryInstanceStatus.request_id == request_id)
    ).all()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    by_sheet_target: dict[str, list] = {}
    for status, detail in rows:
        key = f"{status.dra_type}_{status.instance_label}_{status.entry_type}"
        by_sheet_target.setdefault(key, []).append((status, detail))

    if not by_sheet_target:
        ws = wb.create_sheet("No Data")
        ws["A1"] = "No actionable entries found for this request."
    else:
        for sheet_key, entries in by_sheet_target.items():
            ws = _generate_unique_sheet(wb, sheet_key)

            is_prr = sheet_key.endswith("_PRR")
            base_headers = PRR_DUMP_HEADERS if is_prr else RBAR_DUMP_HEADERS
            audit_headers = ["Action/Decision", "Dependency Note", "Depends On", "Impl Status"]
            _style_header(ws, base_headers + audit_headers)

            num_cols_to_format = [] if is_prr else (4, 5, 9)

            for row_idx, (status, detail) in enumerate(entries, 2):
                row_source = "PENDING_ADD" if status.decision in ADD_DECISIONS else "PENDING_DELETE"
                c_font, c_fill = (MASTER_ADD_FONT, MASTER_ADD_FILL) if row_source == "PENDING_ADD" else (MASTER_DEL_FONT, MASTER_DEL_FILL)

                if is_prr:
                    dump_vals = _unpack_prr_vals(detail)
                else:
                    dump_vals = _unpack_rbar_vals(detail)

                # ── Issue 1 Fix: Pull live data blocks from EntryInstanceRelationship table ──
                rel_records = db.execute(
                    select(EntryInstanceRelationship).where(EntryInstanceRelationship.instance_status_id == status.id)
                ).scalars().all()
                
                if rel_records:
                    depends_on_str = ", ".join([f"{r.related_request_id[:8]}({r.relationship_type})" for r in rel_records])
                else:
                    depends_on_str = "None"

                audit_vals = [status.decision, status.dependency_note or "", depends_on_str, status.impl_status or "N/A"]
                total_vals = dump_vals + audit_vals

                for col_idx, val in enumerate(total_vals, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    cell.font = c_font
                    if c_fill:
                        cell.fill = c_fill
                    
                    if col_idx in num_cols_to_format and isinstance(val, int):
                        cell.number_format = "0"

            _auto_width(ws)

    # ── Integrated Execution Logs Tab ──────────────────────────────────────
    log_ws = wb.create_sheet("Execution_Logs")
    _style_header(log_ws, ["Timestamp (IST)", "Level", "Scope Type", "Message"])

    logs = db.execute(
        select(AuditLog).where(AuditLog.request_id == request_id).order_by(AuditLog.created_at.asc())
    ).scalars().all()

    for r_idx, log in enumerate(logs, 2):
        log_vals = [_ist_format(log.created_at), log.level, log.entry_type, log.message]
        for col, val in enumerate(log_vals, 1):
            log_ws.cell(row=r_idx, column=col, value=val)
    _auto_width(log_ws)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    
    logger.info("Generated Delta workbook output file destination trace: %s", out_path)
    return out_path


def master_odd_excel(db: Session, dra_type: str, instance_label: str, out_path: str) -> str:
    """
    Generate the Master ODD Excel simulation for one specific DRA instance.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ───────────────────────────────────────────────────────────────────────
    # 1. PRR Sheet
    # ───────────────────────────────────────────────────────────────────────
    prr_ws = wb.create_sheet("PRR_PeerRouteRule")
    _style_header(prr_ws, PRR_DUMP_HEADERS + ["Source"])

    prr_snap = db.execute(
        select(DumpSnapshot).where(
            and_(DumpSnapshot.dra_type == dra_type, DumpSnapshot.instance_label == instance_label, DumpSnapshot.object_type == "PRR")
        ).order_by(DumpSnapshot.created_at.desc()).limit(1)
    ).scalars().first()

    base_prr_rows = []
    if prr_snap:
        rows = db.execute(select(PrrDumpRow).where(PrrDumpRow.snapshot_id == prr_snap.id)).scalars().all()
        for r in rows:
            base_prr_rows.append({
                "vals": [
                    "Diameter", "PeerRouteRule", r.name, r.priority, r.param_1, r.cond_operator_1, r.value_1,
                    r.param_2, r.cond_operator_2, r.value_2, r.param_3, r.cond_operator_3, r.value_3, r.param_4,
                    r.cond_operator_4, r.value_4, r.param_5, r.cond_operator_5, r.value_5, r.param_6, r.cond_operator_6,
                    r.value_6, r.action, r.route_list_name, r.diam_ans_code, r.error_message, r.msg_priority,
                    r.msg_cpy_cfg_set, r.vendor_id, r.target_prt_name, r.peer_route_table
                ],
                "realm_key": (r.realm or "").lower(), "meta_source": "DUMP"
            })

    pending_prr = db.execute(
        select(EntryInstanceStatus, EntryInstanceDetail)
        .join(EntryInstanceDetail, EntryInstanceDetail.instance_status_id == EntryInstanceStatus.id)
        .where(and_(
            EntryInstanceStatus.dra_type == dra_type, EntryInstanceStatus.instance_label == instance_label,
            EntryInstanceStatus.entry_type == "PRR",
            EntryInstanceStatus.impl_status.in_([ImplStatus.PENDING.value, ImplStatus.AWAITING_IMPLEMENTATION.value]),
        )).order_by(EntryInstanceStatus.created_at.asc())
    ).all()

    del_prr_realms = {d.realm.lower() for s, d in pending_prr if d.realm and s.decision in DEL_DECISIONS}

    for row in base_prr_rows:
        if row["realm_key"] in del_prr_realms:
            row["meta_source"] = "PENDING_DELETE"

    pending_add_prr_rows = []
    for s, d in pending_prr:
        if s.decision in ADD_DECISIONS:
            pending_add_prr_rows.append({
                "vals": _unpack_prr_vals(d),
                "realm_key": (d.realm or "").lower(), "meta_source": "PENDING_ADD"
            })

    total_ordered_prr = base_prr_rows + pending_add_prr_rows

    for row_idx, r in enumerate(total_ordered_prr, 2):
        row_source = r["meta_source"]
        c_font, c_fill = (MASTER_ADD_FONT, MASTER_ADD_FILL) if row_source == "PENDING_ADD" else (
            (MASTER_DEL_FONT, MASTER_DEL_FILL) if row_source == "PENDING_DELETE" else (MASTER_DUMP_FONT, None)
        )
        total_vals = r["vals"] + [row_source]
        for col, val in enumerate(total_vals, 1):
            cell = prr_ws.cell(row=row_idx, column=col, value=val)
            cell.font = c_font
            if c_fill:
                cell.fill = c_fill
    _auto_width(prr_ws)

    # ───────────────────────────────────────────────────────────────────────
    # 2. RBAR Sheet
    # ───────────────────────────────────────────────────────────────────────
    rbar_ws = wb.create_sheet("RBAR_AddressRange")
    _style_header(rbar_ws, RBAR_DUMP_HEADERS + ["Source"])

    rbar_snap = db.execute(
        select(DumpSnapshot).where(
            and_(DumpSnapshot.dra_type == dra_type, DumpSnapshot.instance_label == instance_label, DumpSnapshot.object_type == "RBAR")
        ).order_by(DumpSnapshot.created_at.desc()).limit(1)
    ).scalars().first()

    base_rbar_rows = []
    if rbar_snap:
        rows = db.execute(select(RbarDumpRow).where(RbarDumpRow.snapshot_id == rbar_snap.id)).scalars().all()
        for r in rows:
            if r.start_addr is not None:
                base_rbar_rows.append({
                    "vals": [
                        "Rbar", "AddressRange", r.table_name,
                        int(r.start_addr), int(r.end_addr) if r.end_addr is not None else None,
                        r.destination, r.pfx_length, r.old_table_name,
                        int(r.old_start_addr) if r.old_start_addr is not None else None, r.old_pfx_length
                    ],
                    "range_key": (int(r.start_addr), int(r.end_addr), (r.destination or "").strip().lower()),
                    "meta_source": "DUMP"
                })

    pending_rbar = db.execute(
        select(EntryInstanceStatus, EntryInstanceDetail)
        .join(EntryInstanceDetail, EntryInstanceDetail.instance_status_id == EntryInstanceStatus.id)
        .where(and_(
            EntryInstanceStatus.dra_type == dra_type, EntryInstanceStatus.instance_label == instance_label,
            EntryInstanceStatus.entry_type == "RBAR",
            EntryInstanceStatus.impl_status.in_([ImplStatus.PENDING.value, ImplStatus.AWAITING_IMPLEMENTATION.value]),
        )).order_by(EntryInstanceStatus.created_at.asc())
    ).all()

    del_rbar_ranges = {
        (int(d.start_addr), int(d.end_addr), (d.destination or "").strip().lower())
        for s, d in pending_rbar if d.start_addr is not None and d.end_addr is not None and s.decision in DEL_DECISIONS
    }

    for row in base_rbar_rows:
        if row["range_key"] in del_rbar_ranges:
            row["meta_source"] = "PENDING_DELETE"

    pending_add_rbar_rows = []
    for s, d in pending_rbar:
        if s.decision in ADD_DECISIONS and d.start_addr is not None and d.end_addr is not None:
            pending_add_rbar_rows.append({
                "vals": _unpack_rbar_vals(d),
                "range_key": (int(d.start_addr), int(d.end_addr), (d.destination or "").strip().lower()),
                "meta_source": "PENDING_ADD"
            })

    total_ordered_rbar = base_rbar_rows + pending_add_rbar_rows

    for row_idx, r in enumerate(total_ordered_rbar, 2):
        row_source = r["meta_source"]
        c_font, c_fill = (MASTER_ADD_FONT, MASTER_ADD_FILL) if row_source == "PENDING_ADD" else (
            (MASTER_DEL_FONT, MASTER_DEL_FILL) if row_source == "PENDING_DELETE" else (MASTER_DUMP_FONT, None)
        )
        total_vals = r["vals"] + [row_source]
        for col, val in enumerate(total_vals, 1):
            cell = rbar_ws.cell(row=row_idx, column=col, value=val)
            cell.font = c_font
            if c_fill:
                cell.fill = c_fill
            
            if col in (4, 5, 9) and isinstance(val, int):
                cell.number_format = "0"
                
    _auto_width(rbar_ws)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    
    logger.info("Generated Master workbook output file destination trace: %s", out_path)
    return out_path